"""
Generate realistic sample resumes (.txt, .docx) and job descriptions for testing.
"""

import os
from docx import Document


def create_samples():
    os.makedirs("examples/resumes", exist_ok=True)
    os.makedirs("examples/job_descriptions", exist_ok=True)

    # 1. Senior Backend Python Resume (TXT)
    backend_resume = """Alex Morgan - Senior Backend Engineer
Email: alex.morgan@example.com | San Francisco, CA

PROFESSIONAL SUMMARY
Senior Software Engineer with 6+ years of experience specializing in high-throughput backend services, distributed systems, and cloud infrastructure using Python, FastAPI, and PostgreSQL.

CORE SKILLS
- Programming: Python, Go, SQL, Bash
- Frameworks: FastAPI, Django, Flask
- Cloud & DevOps: Docker, Kubernetes, AWS (EC2, S3, RDS), Terraform, GitHub Actions, CI/CD, Linux
- Databases & Caching: PostgreSQL, Redis, DynamoDB
- Methodologies: Microservices, REST APIs, System Design, Unit Testing, Agile

EXPERIENCE
Staff Backend Engineer | FinTech Cloud (2022 - Present)
- Architected and scaled financial transaction microservices in Python with FastAPI handling 15,000 requests/sec.
- Containerized legacy monolithic services with Docker and orchestrated deployment onto AWS EKS using Kubernetes.
- Optimized PostgreSQL database queries and implemented Redis caching, reducing p99 response latency by 45%.
- Established automated CI/CD deployment pipelines using GitHub Actions with 95% test coverage.

Software Engineer | SaaS Innovations (2019 - 2022)
- Built customer-facing REST APIs using Django and PostgreSQL.
- Automated AWS cloud provisioning using Terraform scripts.
"""
    with open("examples/resumes/senior_backend_python.txt", "w", encoding="utf-8") as f:
        f.write(backend_resume.strip())

    # 2. Machine Learning Engineer Resume (DOCX)
    doc = Document()
    doc.add_heading("Dr. Elena Rostova - Lead ML Engineer", level=1)
    doc.add_paragraph("Email: elena.ml@example.com | Boston, MA")

    doc.add_heading("Professional Summary", level=2)
    doc.add_paragraph(
        "AI/ML specialist with 5 years of experience developing deep learning architectures, "
        "LLMs, and natural language processing pipelines in Python, PyTorch, and Scikit-Learn."
    )

    doc.add_heading("Technical Competencies", level=2)
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "AI & ML Frameworks"
    table.rows[0].cells[1].text = "PyTorch, TensorFlow, Scikit-Learn, Hugging Face, Keras"
    table.rows[1].cells[0].text = "Core Competencies"
    table.rows[1].cells[1].text = "Machine Learning, Deep Learning, NLP, Computer Vision, LLMs"
    table.rows[2].cells[0].text = "Cloud & Data"
    table.rows[2].cells[1].text = "AWS, Docker, PostgreSQL, Vector Databases, Data Pipelines"

    doc.add_heading("Work Experience", level=2)
    doc.add_paragraph(
        "Lead Machine Learning Engineer | NeuralTech AI (2021 - Present)\n"
        "- Trained and fine-tuned large language models (LLMs) and NLP transformer architectures using PyTorch.\n"
        "- Deployed real-time inference microservices with Docker and Triton on AWS GPU instances.\n"
        "- Built automated ETL data pipelines in Python for multimodal dataset ingestion."
    )
    doc.save("examples/resumes/machine_learning_engineer.docx")

    # 3. Frontend React Developer Resume (TXT) - Designed to test a mismatch against backend JDs
    frontend_resume = """Jordan Lee - Frontend UI/UX Developer
Email: jordan.ui@example.com | New York, NY

SUMMARY
Creative Frontend Engineer with 4 years specializing in responsive user interfaces, React, TypeScript, Next.js, and modern CSS architecture.

TECHNICAL SKILLS
- Languages: JavaScript, TypeScript, HTML, CSS
- Frameworks: React, Next.js, Vue, Tailwind CSS
- Tooling: Webpack, Vite, Jest, Git, Figma

EXPERIENCE
Frontend Developer | WebCraft Studio (2021 - Present)
- Developed enterprise web applications using React, Next.js, and TypeScript.
- Implemented state management with Redux and React Context.
- Built reusable component libraries with Tailwind CSS and Storybook.
"""
    with open("examples/resumes/frontend_react_dev.txt", "w", encoding="utf-8") as f:
        f.write(frontend_resume.strip())

    # 4. Job Description 1: Senior Python Backend Engineer
    backend_jd = """Job Title: Senior Backend Engineer - Python & Cloud Infrastructure
Company: Apex Cloud Systems
Location: Remote / Hybrid

About the Role:
We are seeking a Senior Backend Engineer to design, build, and scale our core cloud microservices. You will take ownership of backend APIs, database architecture, and deployment pipelines.

Key Responsibilities:
- Design and maintain high-performance REST APIs in Python using FastAPI or Django.
- Manage and optimize relational databases (PostgreSQL) and caching layers (Redis).
- Build containerized applications using Docker and deploy them on AWS cloud infrastructure with Kubernetes.
- Drive automated CI/CD pipeline improvements and maintain test-driven development standards.

Required Qualifications:
- 4+ years software engineering experience in Python.
- Demonstrated hands-on expertise with Docker, Kubernetes, and AWS.
- Strong proficiency with PostgreSQL and Redis.
- Solid background in Microservices, System Design, and CI/CD automation.
"""
    with open("examples/job_descriptions/senior_python_backend_jd.txt", "w", encoding="utf-8") as f:
        f.write(backend_jd.strip())

    # 5. Job Description 2: Lead Machine Learning Engineer
    ml_jd = """Job Title: Staff / Lead Machine Learning Engineer
Company: Cognitive Scale AI
Location: Remote

Role Overview:
Cognitive Scale is looking for an experienced Machine Learning Engineer to spearhead our generative AI and NLP initiative.

Requirements:
- 4+ years experience in Python, PyTorch, and Scikit-Learn.
- Proven track record with Natural Language Processing (NLP), Deep Learning, and LLMs.
- Experience with Docker containerization and deploying models to AWS.
- Knowledge of Vector Databases and Data Engineering pipelines.
"""
    with open("examples/job_descriptions/lead_ml_engineer_jd.txt", "w", encoding="utf-8") as f:
        f.write(ml_jd.strip())

    print("Successfully generated sample resumes and job descriptions in 'examples/'.")


if __name__ == "__main__":
    create_samples()