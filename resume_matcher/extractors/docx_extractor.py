"""
DOCX document extractor using python-docx.
"""

import io
from typing import Union
from docx import Document
from resume_matcher.extractors.base import BaseExtractor, ExtractionError


class DocxExtractor(BaseExtractor):
    """Extractor for Microsoft Word (.docx) documents."""

    def extract_from_file(self, file_path: str) -> str:
        """Extract text from a DOCX file path."""
        self._validate_file_path(file_path)
        try:
            doc = Document(file_path)
            return self._extract_text_from_doc(doc)
        except Exception as e:
            raise ExtractionError(f"Failed to read DOCX file '{file_path}': {e}") from e

    def extract_from_bytes(self, file_bytes: Union[bytes, io.BytesIO], filename: str = "") -> str:
        """Extract text from DOCX in-memory bytes."""
        try:
            if isinstance(file_bytes, bytes):
                stream = io.BytesIO(file_bytes)
            elif isinstance(file_bytes, io.BytesIO):
                stream = file_bytes
            else:
                raise ExtractionError(f"Unsupported byte input type: {type(file_bytes)}")

            doc = Document(stream)
            return self._extract_text_from_doc(doc)
        except Exception as e:
            raise ExtractionError(f"Failed to extract DOCX bytes: {e}") from e

    @staticmethod
    def _extract_text_from_doc(doc: Document) -> str:
        """Extract text from both document paragraphs and structured tables."""
        parts = []

        # 1. Paragraphs
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)

        # 2. Tables (resumes often put skills/contact info in tables)
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    # Remove duplicate adjacent cells caused by merged table columns
                    deduped = []
                    for c in row_cells:
                        if not deduped or c != deduped[-1]:
                            deduped.append(c)
                    parts.append(" | ".join(deduped))

        full_text = "\n".join(parts)
        full_text = full_text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
        return full_text.strip()