import unittest
import tempfile
import os
import io
from docx import Document

from resume_matcher.extractors import (
    TxtExtractor,
    PdfExtractor,
    DocxExtractor,
    DocumentExtractorFactory,
    ExtractionError,
)


class TestExtractors(unittest.TestCase):

    def test_txt_extractor(self):
        extractor = TxtExtractor()
        text = "Senior Python Developer\nDjango, FastAPI, PostgreSQL"
        result = extractor.extract_from_bytes(text.encode("utf-8"))
        self.assertEqual(result, text)

    def test_docx_extractor(self):
        # Create an in-memory DOCX
        doc = Document()
        doc.add_heading("John Doe - Resume", level=1)
        doc.add_paragraph("Skills: Python, Docker, Kubernetes")
        
        # Add table
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Education"
        table.rows[0].cells[1].text = "B.S. Computer Science"

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)

        extractor = DocxExtractor()
        extracted = extractor.extract_from_bytes(buf)
        self.assertIn("John Doe - Resume", extracted)
        self.assertIn("Skills: Python, Docker, Kubernetes", extracted)
        self.assertIn("Education | B.S. Computer Science", extracted)

    def test_factory_resolution(self):
        txt_ext = DocumentExtractorFactory.get_extractor("resume.txt")
        self.assertIsInstance(txt_ext, TxtExtractor)

        pdf_ext = DocumentExtractorFactory.get_extractor("resume.pdf")
        self.assertIsInstance(pdf_ext, PdfExtractor)

        docx_ext = DocumentExtractorFactory.get_extractor("resume.docx")
        self.assertIsInstance(docx_ext, DocxExtractor)

    def test_factory_unsupported_format(self):
        with self.assertRaises(ExtractionError):
            DocumentExtractorFactory.get_extractor("resume.xyz")


if __name__ == "__main__":
    unittest.main()